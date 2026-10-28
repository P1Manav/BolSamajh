"""Document loader — PDF, DOCX, and web page ingestion with encryption detection, OCR bounds, and HTML title extraction."""

import html
import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

log = logging.getLogger(__name__)

# Limits for OCR processing to prevent resource exhaustion
MAX_OCR_PAGES = 30
MAX_OCR_FILE_SIZE_BYTES = 50 * 1024 * 1024  # 50 MB


@dataclass
class DocumentPage:
    """A single page/section of a loaded document."""

    text: str
    page_number: int = 0  # 1-indexed (0 = unknown)
    section: str = ""  # Section heading if available
    source: str = ""  # File path or URL


@dataclass
class Document:
    """A loaded document with page-level content."""

    title: str
    source: str  # File path or URL
    pages: list[DocumentPage] = field(default_factory=list)
    doc_type: str = ""  # "pdf", "docx", "web", "txt", "md"

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())

    @property
    def page_count(self) -> int:
        return len(self.pages)


def _extract_page_ocr(page, page_num: int, filename: str, file_size: int = 0) -> str:
    """Attempt local OCR on scanned PDF page within bounded size/page constraints."""
    if page_num > MAX_OCR_PAGES:
        log.warning(
            "Page %d of '%s' exceeds OCR limit (%d pages); skipping OCR fallback.",
            page_num,
            filename,
            MAX_OCR_PAGES,
        )
        return ""

    if file_size > MAX_OCR_FILE_SIZE_BYTES:
        log.warning(
            "PDF '%s' size (%d bytes) exceeds OCR limit (%d bytes); skipping OCR fallback.",
            filename,
            file_size,
            MAX_OCR_FILE_SIZE_BYTES,
        )
        return ""

    # 1. Try PyMuPDF native OCR bindings
    try:
        tp = page.get_textpage_ocr(language="hin+eng")
        ocr_text = tp.extractText().strip()
        if ocr_text:
            log.info(
                "OCR successfully extracted %d characters from page %d of '%s'",
                len(ocr_text),
                page_num,
                filename,
            )
            return ocr_text
    except Exception as e:
        log.debug("PyMuPDF get_textpage_ocr failed on page %d: %s", page_num, e)

    # 2. Try pytesseract if installed locally
    try:
        import pytesseract
        from PIL import Image

        pix = page.get_pixmap(dpi=150)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        ocr_text = pytesseract.image_to_string(img, lang="hin+eng").strip()
        if ocr_text:
            log.info(
                "pytesseract extracted %d characters from page %d of '%s'",
                len(ocr_text),
                page_num,
                filename,
            )
            return ocr_text
    except ImportError:
        log.warning(
            "Page %d of '%s' appears scanned but OCR failed: 'pytesseract' is not installed. "
            "To enable local OCR for scanned PDFs, install pytesseract and Tesseract engine: "
            "Windows: `winget install UB-Mannheim.TesseractOCR` && `pip install pytesseract` | "
            "Linux: `sudo apt install tesseract-ocr tesseract-ocr-hin`",
            page_num,
            filename,
        )
    except Exception as e:
        log.warning("Tesseract OCR failed on page %d of '%s': %s.", page_num, filename, e)

    return ""


def load_pdf(file_path: str) -> Document:
    """Load a PDF file using PyMuPDF with encryption detection and bounded OCR fallback."""
    import fitz  # PyMuPDF

    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"PDF not found: {file_path}")

    log.info("Loading PDF: %s", path.name)
    file_size = path.stat().st_size
    doc = fitz.open(str(path))

    # Detect encrypted or password-protected PDF
    if doc.is_encrypted or doc.needs_pass:
        doc.close()
        raise ValueError(
            f"PDF '{path.name}' is encrypted or password-protected. "
            "Please remove password protection or decrypt the file before loading."
        )

    pages = []
    blank_pages = []
    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text("text").strip()

        if not text:
            log.warning(
                "Page %d of '%s' has no embedded text (possible scan); attempting bounded OCR fallback...",
                page_num + 1,
                path.name,
            )
            text = _extract_page_ocr(page, page_num + 1, path.name, file_size=file_size)

        if text:
            pages.append(
                DocumentPage(
                    text=text,
                    page_number=page_num + 1,
                    source=str(path),
                )
            )
        else:
            blank_pages.append(page_num + 1)
            log.warning("Page %d of '%s' contains no readable text.", page_num + 1, path.name)

    if blank_pages:
        log.warning(
            "PDF '%s' has %d page(s) without extractable text: %s",
            path.name,
            len(blank_pages),
            blank_pages,
        )

    title = doc.metadata.get("title", "") or path.stem
    doc.close()

    if not pages:
        raise ValueError(
            f"PDF '{path.name}' contains no readable text. It appears to be scanned images "
            "and local OCR was unable to extract text. Please ensure Tesseract is installed for image PDFs."
        )

    result = Document(
        title=title,
        source=str(path),
        pages=pages,
        doc_type="pdf",
    )
    log.info("Loaded PDF: %s (%d pages with text)", title, len(pages))
    return result


def load_docx(file_path: str) -> Document:
    """Load a Word document using python-docx, including tables and checking for empty content."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX not found: {file_path}")

    if path.suffix.lower() == ".doc":
        raise ValueError(
            f"Unsupported format '{path.suffix}': Legacy Word (.doc) files are not supported. "
            "Please convert the file to modern .docx or .pdf."
        )

    from docx import Document as DocxDoc

    log.info("Loading DOCX: %s", path.name)
    try:
        docx = DocxDoc(str(path))
    except Exception as e:
        raise ValueError(f"Failed to read DOCX file '{path.name}': {e}")

    pages = []
    current_section = ""
    current_text_parts = []
    page_num = 1

    # Extract paragraphs
    for para in docx.paragraphs:
        if para.style and para.style.name and para.style.name.startswith("Heading"):
            if current_text_parts:
                pages.append(
                    DocumentPage(
                        text="\n".join(current_text_parts),
                        page_number=page_num,
                        section=current_section,
                        source=str(path),
                    )
                )
                page_num += 1
                current_text_parts = []
            current_section = para.text.strip()

        if para.text.strip():
            current_text_parts.append(para.text.strip())

    # Include DOCX tables
    for table in docx.tables:
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                table_rows.append(" | ".join(row_cells))
        if table_rows:
            current_text_parts.append("\n".join(table_rows))

    # Save last section
    if current_text_parts:
        pages.append(
            DocumentPage(
                text="\n".join(current_text_parts),
                page_number=page_num,
                section=current_section,
                source=str(path),
            )
        )

    if not pages or not any(p.text.strip() for p in pages):
        raise ValueError(f"DOCX '{path.name}' is empty or contains no readable text.")

    title = path.stem
    result = Document(
        title=title,
        source=str(path),
        pages=pages,
        doc_type="docx",
    )
    log.info("Loaded DOCX: %s (%d sections)", title, len(pages))
    return result


def load_webpage(url: str, html_content: str | None = None) -> Document:
    """Load a web page using trafilatura and extract the HTML <title>."""
    import trafilatura

    log.info("Loading web page: %s", url[:80])
    downloaded = html_content if html_content is not None else trafilatura.fetch_url(url)
    if not downloaded:
        raise ValueError(f"Could not fetch URL: {url}")

    # Extract HTML <title>
    title = ""
    title_match = re.search(r"<title[^>]*>(.*?)</title>", downloaded, re.IGNORECASE | re.DOTALL)
    if title_match:
        raw_title = title_match.group(1).strip()
        title = html.unescape(raw_title)

    text = trafilatura.extract(
        downloaded,
        include_comments=False,
        include_tables=True,
        include_links=False,
    )

    if not text:
        raise ValueError(f"Could not extract text from: {url}")

    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    if len(paragraphs) <= 1:
        paragraphs = [p.strip() for p in text.splitlines() if p.strip()]

    pages = []
    for i, para in enumerate(paragraphs, 1):
        pages.append(
            DocumentPage(
                text=para,
                page_number=i,
                section=f"Section {i}",
                source=url,
            )
        )

    if not title:
        title = url.split("/")[-1] or "Web Page"

    result = Document(
        title=title,
        source=url,
        pages=pages,
        doc_type="web",
    )
    log.info("Loaded web page '%s': %d sections extracted", title, len(pages))
    return result


def load_text_file(path_str: str) -> Document:
    """Load plain text or markdown file, splitting on both CRLF and LF."""
    path = Path(path_str)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path_str}")

    content = path.read_text(encoding="utf-8", errors="replace")
    # Normalize CRLF and CR to LF
    normalized = content.replace("\r\n", "\n").replace("\r", "\n")
    paragraphs = [p.strip() for p in normalized.split("\n\n") if p.strip()]
    if not paragraphs:
        paragraphs = [normalized.strip()] if normalized.strip() else ["(Empty document)"]

    pages = []
    for i, para in enumerate(paragraphs, 1):
        pages.append(
            DocumentPage(
                text=para,
                page_number=i,
                section=f"Section {i}",
                source=str(path),
            )
        )

    title = path.stem
    result = Document(
        title=title,
        source=str(path),
        pages=pages,
        doc_type=path.suffix.lstrip(".").lower() or "txt",
    )
    log.info("Loaded text file: %s (%d sections)", title, len(pages))
    return result


def load_document(source: str) -> Document:
    """Auto-detect and load a document from file path or URL."""
    source = source.strip()

    if source.startswith(("http://", "https://")):
        return load_webpage(source)

    path = Path(source)
    suffix = path.suffix.lower()

    if suffix == ".doc":
        raise ValueError(
            f"Unsupported format '{suffix}': Legacy Word (.doc) files are not supported. "
            "Please convert the file to modern .docx or .pdf."
        )
    elif suffix == ".pdf":
        return load_pdf(source)
    elif suffix == ".docx":
        return load_docx(source)
    elif suffix in (".txt", ".md", ".markdown", ".rst"):
        return load_text_file(source)
    else:
        raise ValueError(
            f"Unsupported document type: '{suffix}'. Supported: .pdf, .docx, .txt, .md, or a URL."
        )
